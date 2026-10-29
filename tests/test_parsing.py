import io

import pytest
from docx import Document as Docx
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

from enterprise_rag.ingestion.parsing import UnparseableDocument, make_docling_parse


@pytest.fixture(scope="module")
def parse():
    return make_docling_parse()  # builds the converter once


def docx_bytes() -> bytes:
    d = Docx()
    d.add_paragraph("Applies to all stores.")
    d.add_heading("Refund Policy", 1)
    d.add_paragraph("Refunds within 30 days.")
    d.add_heading("Exceptions", 2)
    d.add_paragraph("Gift cards are final.")
    table = d.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "Item", (0, 1): "Days", (1, 0): "Shoes", (1, 1): "30"}.items():
        table.cell(r, c).text = text
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pdf_bytes() -> bytes:
    styles, buf = getSampleStyleSheet(), io.BytesIO()
    SimpleDocTemplate(buf).build(
        [
            Paragraph("Refund Policy", styles["Heading1"]),
            Paragraph("Refunds within 30 days.", styles["BodyText"]),
            Paragraph("Exceptions", styles["Heading2"]),
            Paragraph("Gift cards are final.", styles["BodyText"]),
        ]
    )
    return buf.getvalue()


# Break caught: headings dropped or flattened, so citations cannot say where a chunk came
# from; text before the first heading lost; tables discarded.
def test_docx_sections_follow_the_heading_tree(parse):
    sections = parse(docx_bytes(), "original.docx")

    assert [(s.heading_path, s.page) for s in sections] == [
        ("", None),
        ("Refund Policy", None),
        ("Refund Policy > Exceptions", None),
    ]
    assert sections[0].text == "Applies to all stores."
    assert sections[1].text == "Refunds within 30 days."
    assert sections[2].text.startswith("Gift cards are final.\n")
    assert "| Shoes" in sections[2].text  # table kept, as markdown


# Break caught: section offsets that do not line up with the document text the chunker and
# citations index into.
def test_offsets_lay_sections_end_to_end(parse):
    sections = parse(docx_bytes(), "original.docx")

    assert [s.offset for s in sections] == [0, 24, 49]


# Break caught: PDF pages lost, so citations cannot point at a page.
def test_pdf_sections_carry_page_numbers(parse):
    sections = parse(pdf_bytes(), "original.pdf")

    # Docling reports every PDF heading at level 1, so PDF paths are one heading deep.
    assert [(s.heading_path, s.text, s.page) for s in sections] == [
        ("Refund Policy", "Refunds within 30 days.", 1),
        ("Exceptions", "Gift cards are final.", 1),
    ]


def pdf_with_pages(n: int) -> bytes:
    styles, buf, story = getSampleStyleSheet(), io.BytesIO(), []
    for i in range(n):
        story += [Paragraph(f"Page {i + 1}", styles["BodyText"]), PageBreak()]
    SimpleDocTemplate(buf).build(story)
    return buf.getvalue()


# Break caught: every format Docling knows (HTML, XML, images, ...) reachable from an upload,
# widening the parser attack surface beyond what we support.
def test_unsupported_formats_are_refused(parse):
    with pytest.raises(UnparseableDocument):
        parse(b"<html><body><h1>Hi</h1></body></html>", "page.html")


# Break caught: a huge PDF tying up a worker for minutes and exhausting memory.
def test_pdfs_over_the_page_cap_are_refused():
    with pytest.raises(UnparseableDocument, match="pages"):
        make_docling_parse(max_pages=2)(pdf_with_pages(3), "original.pdf")


# Break caught: parser crashes on malformed files escaping as arbitrary errors, so the worker
# cannot tell "bad document, stop" from "transient, retry".
def test_corrupt_files_raise_unparseable_document(parse):
    with pytest.raises(UnparseableDocument):
        parse(b"%PDF-1.7 truncated garbage", "original.pdf")
