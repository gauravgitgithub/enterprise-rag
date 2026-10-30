"""Upload the eval corpus to the local stack as one tenant's DOCX files and wait for ingestion.

Goes through the deployed pipeline: SeaweedFS, register_document, outbox, relay, SQS, worker.
Prints the tenant id for load_test.py and chaos.py. Idempotent per tenant name.
"""

import asyncio
import hashlib
import io
import re
import sys
import time
import uuid
from pathlib import Path

import boto3
import docx
from devstack import GROUP, PG_HOST, S3_URL, secrets
from sqlalchemy import insert, select

from enterprise_rag.constants.documents import DOCUMENT_STATUS_FAILED, DOCUMENT_STATUS_READY
from enterprise_rag.context import RequestContext
from enterprise_rag.db.models import Document, Tenant
from enterprise_rag.db.session import make_engine, make_session_factory
from enterprise_rag.ingestion.service import register_document
from enterprise_rag.repositories.uow import UnitOfWork

BUCKET = "documents"
TENANT_NAME = sys.argv[1] if len(sys.argv) > 1 else "load-test"
TIMEOUT_S = 1800
CORPUS = sorted(Path("evals/corpus").glob("*.md"))


def to_docx(markdown: str) -> bytes:
    document = docx.Document()
    for line in markdown.splitlines():
        if heading := re.match(r"^(#+) (.+)", line):
            document.add_heading(heading.group(2), level=len(heading.group(1)))
        elif line.strip():
            document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


async def main() -> None:
    env = secrets()
    admin = make_engine(f"postgresql+asyncpg://postgres:{env['POSTGRES_PASSWORD']}@{PG_HOST}/rag")
    async with admin.begin() as conn:
        tenant = (
            await conn.execute(select(Tenant.id).where(Tenant.name == TENANT_NAME))
        ).scalar() or (
            await conn.execute(insert(Tenant).values(name=TENANT_NAME).returning(Tenant.id))
        ).scalar_one()
    await admin.dispose()

    app = make_session_factory(
        make_engine(f"postgresql+asyncpg://rag_app:{env['RAG_APP_PASSWORD']}@{PG_HOST}/rag")
    )
    ctx = RequestContext(tenant, frozenset({GROUP}))
    s3 = boto3.client(
        "s3",
        endpoint_url=S3_URL,
        aws_access_key_id="local",
        aws_secret_access_key="local",  # noqa: S106  SeaweedFS accepts any credentials
        region_name="eu-west-1",
    )
    ids = []
    for path in CORPUS:
        raw = to_docx(path.read_text())
        key = f"tenants/{tenant}/docs/{uuid.uuid4()}/v1/{path.stem}.docx"
        s3.put_object(Bucket=BUCKET, Key=key, Body=raw)
        async with UnitOfWork(app, ctx) as uow:
            doc, _ = await register_document(
                uow,
                Document(
                    title=path.stem,
                    source_uri=f"s3://{BUCKET}/{key}",
                    content_hash=hashlib.sha256(raw).hexdigest(),
                    acl_groups=[GROUP],
                ),
            )
            ids.append(doc.id)

    started = time.monotonic()
    while time.monotonic() - started < TIMEOUT_S:
        async with UnitOfWork(app, ctx) as uow:
            statuses = [(await uow.documents.get(i)).status for i in ids]  # type: ignore[union-attr]
        ready = statuses.count(DOCUMENT_STATUS_READY)
        print(f"{ready}/{len(ids)} ready after {time.monotonic() - started:.0f}s", flush=True)
        if ready == len(ids) or DOCUMENT_STATUS_FAILED in statuses:
            break
        await asyncio.sleep(10)
    print(f"tenant {tenant}")


asyncio.run(main())
